#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CUMCM 2026 C题 问题2 —— 最终策略: w1预测(上周同星期) + 余量(γ=1.05, δ=0.98)
流程: 日前(0:00) w1预测→日LP计划(S(144)=S(0)=S0)→日内实际模拟(储能限额内调节,
      缺口经 减充电→加放电→紧急购电5倍; 富余经 减放电→加充电→弃光)。
约定: 每行时刻戳=区间起点, 第d行覆盖 0:10(d)→0:10(d+1); 储电量取实际轨迹。
输出: result2.xlsx (2025.2.1-12.31), summary2.json, q2_data.npz
"""
import json
import numpy as np
import openpyxl
from scipy.optimize import linprog

BASE = "/home/jason/DataDisk/Jason's study/数学建模大赛/题目/CUMCM2026Problems/C题"

c1 = openpyxl.load_workbook(BASE + "/附件/附件1.xlsx", data_only=True)["Sheet1"]
rows1 = list(c1.iter_rows(min_row=2, values_only=True))
price = np.array([float(r[1]) for r in rows1])
L_typ = np.array([float(r[2]) for r in rows1])
PV_typ = np.array([float(r[3]) for r in rows1])

def load_cols(path, sheet):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return [r[0] for r in rows], np.array([[float(v) for v in r[1:145]] for r in rows])

dates, LD = load_cols(BASE + "/附件/附件2.xlsx", "小区负载")
_, PVD = load_cols(BASE + "/附件/附件2.xlsx", "光伏发电实际功率")
NDAY = 365
T, dt = 144, 1.0 / 6.0
ETA_C = ETA_D = 0.9
PMAX, SMIN, SMAX = 5000.0, 1200.0, 10800.0
EMULT = 5.0
ip, ich, idi, iq, is_ = 0, T, 2 * T, 3 * T, 4 * T
nx = 5 * T + 1
feb1 = 31

GAMMA, DELTA = 1.05, 0.98          # 余量系数(扫描选定)

def forecast(d):
    """w1: 上周同星期实际; 首周回退持续性; 1月1日用附件1典型日"""
    if d == 0: return L_typ, PV_typ
    if d >= 7: return LD[d-7], PVD[d-7]
    return LD[d-1], PVD[d-1]

def plan_lp(Lhat, PVhat, S0, price_vec=None):
    pr = price if price_vec is None else price_vec
    f = np.zeros(nx); f[ip:ip+T] = pr * dt
    Aeq, beq = [], []
    Lh, Ph = Lhat * GAMMA, PVhat * DELTA
    for t in range(T):
        r = np.zeros(nx)
        r[ip+t] = 1; r[idi+t] = 1; r[ich+t] = -1; r[iq+t] = -1
        Aeq.append(r); beq.append(Lh[t] - Ph[t])
    for t in range(T):
        r = np.zeros(nx)
        r[is_+t+1] = 1; r[is_+t] = -1
        r[ich+t] = -ETA_C*dt; r[idi+t] = dt/ETA_D
        Aeq.append(r); beq.append(0.0)
    for col, val in [(is_, S0), (is_+T, S0)]:
        r = np.zeros(nx); r[col] = 1
        Aeq.append(r); beq.append(val)
    bnd = ([(0, None)]*T + [(0, PMAX)]*2*T + [(0, None)]*T + [(SMIN, SMAX)]*(T+1))
    res = linprog(f, A_eq=np.array(Aeq), b_eq=np.array(beq), bounds=bnd, method="highs")
    assert res.status == 0, f"计划LP失败: {res.message}"
    x = res.x
    return x[ip:ip+T], x[ich:ich+T], x[idi:idi+T], res.fun

def simulate(ps, chs, diss, L, PV, S0):
    e = np.zeros(T); ch_r = np.zeros(T); dis_r = np.zeros(T)
    S = np.empty(T+1); S[0] = S0
    for t in range(T):
        dis_max = min(PMAX, max(S[t]-SMIN, 0.0)*ETA_D/dt)
        ch_max  = min(PMAX, max(SMAX-S[t], 0.0)/ETA_C/dt)
        dis = min(diss[t], dis_max); ch = min(chs[t], ch_max)
        resid = (L[t] - ps[t] - PV[t]) - (dis - ch)
        if resid > 0:
            r = min(ch, resid); ch -= r; resid -= r
            r = min(dis_max - dis, resid); dis += r; resid -= r
            e[t] = max(resid, 0.0)
        else:
            sur = -resid
            r = min(dis, sur); dis -= r; sur -= r
            r = min(ch_max - ch, sur); ch += r; sur -= r
        ch_r[t], dis_r[t] = ch, dis
        S[t+1] = S[t] + ETA_C*ch*dt - dis*dt/ETA_D
    return e, ch_r, dis_r, S

# ---------- 全年滚动 ----------
def run_year(prices=None):
    plans_p = np.zeros((NDAY, T)); plans_ch = np.zeros((NDAY, T)); plans_dis = np.zeros((NDAY, T))
    Emerg = np.zeros((NDAY, T)); ch_r = np.zeros((NDAY, T)); dis_r = np.zeros((NDAY, T))
    Sall = np.zeros((NDAY, T+1)); plan_cost = np.zeros(NDAY)
    S0 = 6000.0
    for d in range(NDAY):
        Lhat, PVhat = forecast(d)
        pv_d = None if prices is None else prices[d]
        ps, chs, diss, pc = plan_lp(Lhat, PVhat, S0, pv_d)
        e, cr, dr, S = simulate(ps, chs, diss, LD[d], PVD[d], S0)
        plans_p[d], plans_ch[d], plans_dis[d] = ps, chs, diss
        Emerg[d], ch_r[d], dis_r[d], Sall[d], plan_cost[d] = e, cr, dr, S, pc
        S0 = S[T]
    return dict(plans_p=plans_p, plans_ch=plans_ch, plans_dis=plans_dis,
                Emerg=Emerg, ch_r=ch_r, dis_r=dis_r, Sall=Sall, plan_cost=plan_cost)

if __name__ == "__main__":
    _R = run_year()
    plans_p, plans_ch, plans_dis = _R["plans_p"], _R["plans_ch"], _R["plans_dis"]
    Emerg, ch_r, dis_r, Sall, plan_cost = _R["Emerg"], _R["ch_r"], _R["dis_r"], _R["Sall"], _R["plan_cost"]

    rep = slice(feb1, NDAY)
    E_p = plans_p[rep].sum()*dt
    C_p = float(plan_cost[rep].sum())
    E_e = Emerg[rep].sum()*dt
    C_e = float((Emerg[rep]*price*EMULT).sum()*dt)
    C_tot = C_p + C_e
    edays = int((Emerg[rep].sum(1) > 1e-9).sum())
    print("========== 问题2 最终策略 (w1预测, γ=1.05, δ=0.98) ==========")
    print(f"报告期 2025.2.1-12.31:")
    print(f"  计划购电量 {E_p:,.1f} kWh, 计划购电费 {C_p:,.2f} 元")
    print(f"  紧急购电量 {E_e:,.1f} kWh, 紧急购电费 {C_e:,.2f} 元")
    print(f"  总费用 {C_tot:,.2f} 元, 紧急购电天数 {edays}/334")
    print(f"  实际储电量范围 [{Sall.min():.1f}, {Sall.max():.1f}] kWh")
    assert Sall.min() >= SMIN-1e-6 and Sall.max() <= SMAX+1e-6
    viol = sum(1 for d in range(NDAY) for t in range(T)
               if Emerg[d, t] > 1e-9 and not (ch_r[d, t] < 1e-9 and
               (dis_r[d, t] > PMAX-1e-6 or Sall[d, t+1] < SMIN+1e-6)))
    print(f"  紧急购电合理性违规: {viol} 个")

    # ---------- 表3指定日期 ----------
    def tag(m):
        plus = "+1" if m >= 1440 else ""
        m %= 1440
        return f"{m//60}:{m%60:02d}{plus}"

    # 表2 钟点日时段块: "0:00-4:00" = 前一日末段(0:00-0:10) + 当日第1-23段
    def bsum(arr, d, k):
        if k == 0:
            return float(arr[d - 1][143] + arr[d][0:23].sum())
        seg = [(23, 47), (47, 71), (71, 95), (95, 119), (119, 143)][k - 1]
        return float(arr[d][seg[0]:seg[1]].sum())
    bnames = ["0:00-4:00","4:00-8:00","8:00-12:00","12:00-16:00","16:00-20:00","20:00-24:00"]
    idx_t1 = [59, 71, 83, 95, 107, 119]
    names_t1 = ["10:00-10:10","12:00-12:10","14:00-14:10",
                "16:00-16:10","18:00-18:10","20:00-20:10"]

    def fmt_date(d):
        x = dates[d]
        return f"{x.year}.{x.month}.{x.day}"

    def emerg_windows(e):
        wins = []; t = 0
        while t < T:
            if e[t] > 1e-9:
                j = t
                while j+1 < T and e[j+1] > 1e-9: j += 1
                wins.append([f"{tag(10*(t+1))}-{tag(10*(j+2))}", float(e[t:j+1].sum()*dt)])
                t = j+1
            else: t += 1
        return wins

    table3 = {}
    for d in range(NDAY):
        if fmt_date(d) in ["2025.3.20", "2025.6.21", "2025.9.23", "2025.12.21"]:
            Ep = plans_p[d]*dt
            Ech, Edis = ch_r[d]*dt, dis_r[d]*dt
            e = Emerg[d]*dt
            wins = emerg_windows(Emerg[d])
            table3[fmt_date(d)] = {
                "plan_t1": {n: float(Ep[i]) for n, i in zip(names_t1, idx_t1)},
                "plan_total": float(Ep.sum()),
                "plan_cost": float(plan_cost[d]),
                "blocks": {n: [bsum(ch_r * dt, d, k), bsum(dis_r * dt, d, k)]
                           for k, n in enumerate(bnames)},
                "s000": float(Sall[d-1][143] if d > 0 else 6000.0),
                "s2400": float(Sall[d][143]),
                "emerg_windows": wins,
                "emerg_total": float(e.sum()),
            }

    print("\n========== 表3 指定日期(最终策略) ==========")
    for k, v in table3.items():
        print(f"\n--- {k} ---")
        print(f"  全天计划购电量 {v['plan_total']:.2f} kWh, 计划购电费 {v['plan_cost']:.2f} 元")
        print("  表1:", {n: round(x, 2) for n, x in v["plan_t1"].items()})
        print("  表2:", {n: [round(a, 2), round(b, 2)] for n, (a, b) in v["blocks"].items()})
        print(f"  s(0:00)={v['s000']:.2f}, s(24:00)={v['s2400']:.2f}")
        print(f"  紧急购电: {[(w, round(x,2)) for w,x in v['emerg_windows']] if v['emerg_windows'] else '无'} (合计 {v['emerg_total']:.2f} kWh)")

    json.dump({"annual": {"E_plan": E_p, "C_plan": C_p, "E_emerg": E_e, "C_emerg": C_e,
                          "C_total": C_tot, "emerg_days": edays},
               "table3": table3},
              open(BASE + "/Q2/summary2.json", "w"), ensure_ascii=False, indent=1)

    # ---------- 写 result2.xlsx ----------
    out = openpyxl.Workbook()
    ws1 = out.active; ws1.title = "计划购电量"
    hdr = ["日期\\时间"] + [f"{tag(10*(i+1))}-{tag(10*(i+2))}" for i in range(143)]
    hdr += ["0:00-0:10+1", "全天购电量", "全天购电费"]
    hdr[42] = "7:0-7:10"          # 保留官方模板第43列的 typo 以逐字一致
    ws1.append(hdr)
    for d in range(feb1, NDAY):
        Ep = plans_p[d]*dt
        ws1.append([dates[d]] + [round(float(x), 4) for x in Ep] +
                   [round(float(Ep.sum()), 4), round(float(plan_cost[d]), 4)])

    ws2 = out.create_sheet("充放电量")
    ws2.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    import datetime as _dt
    for d in range(feb1, NDAY):
        Ech, Edis = ch_r[d]*dt, dis_r[d]*dt
        s000 = Sall[d-1][143] if d > 0 else 6000.0
        s2400 = Sall[d][143]
        for k, n in enumerate(bnames):
            r = [dates[d] if k == 0 else None, n,
                 round(bsum(ch_r * dt, d, k), 4), round(bsum(dis_r * dt, d, k), 4)]
            if k == 0:   r += [_dt.time(0, 0), round(float(s000), 4)]   # 与模板同为 time 对象
            elif k == 1: r += ["24:00", round(float(s2400), 4)]
            else:        r += [None, None]
            ws2.append(r)

    ws3 = out.create_sheet("紧急购电量")
    ws3.append(["日期", "购电时间段", "购电量"])
    for d in range(feb1, NDAY):
        first = True
        for w, energy in emerg_windows(Emerg[d]):
            ws3.append([dates[d] if first else None, w, round(energy, 4)])
            first = False
    out.save(BASE + "/Q2/result2.xlsx")
    print("\n已写出 result2.xlsx")

    np.savez(BASE + "/Q2/q2_data.npz", plans_p=plans_p, plans_ch=plans_ch,
             plans_dis=plans_dis, Emerg=Emerg, Sall=Sall, plan_cost=plan_cost,
             ch_r=ch_r, dis_r=dis_r, allow_pickle=True)
    print("已保存 q2_data.npz")
