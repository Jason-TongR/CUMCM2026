#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CUMCM 2026 C题 问题4 —— 波动电价下重算问题2/问题3
电价: 附件4 逐日曲线(365×144), 假设日前可知(日前市场公布);
     紧急购电=当时实际电价×5, 调整结算=附件4实际价格。
Q4-2 = 问题2最终策略(w1预测, γ=1.05, δ=0.98) 换波动电价 → result4-2.xlsx
Q4-3 = 问题3最终策略(V4F, γ=1.02, δ=0.98) 换波动电价 → result4-3.xlsx
"""
import json
import importlib.util
import numpy as np
import openpyxl

BASE = "/home/jason/DataDisk/Jason's study/数学建模大赛/题目/CUMCM2026Problems/C题"

def imp(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m

q2 = imp(BASE + "/Q2/solve_q2.py", "q2")
q3 = imp(BASE + "/Q3/solve_q3.py", "q3")

# ---------- 附件4 波动电价 ----------
ws4 = openpyxl.load_workbook(BASE + "/附件/附件4.xlsx", data_only=True)["Sheet1"]
rows4 = list(ws4.iter_rows(min_row=2, values_only=True))
dates4 = [r[0] for r in rows4]
PM = np.array([[float(v) for v in r[1:145]] for r in rows4])   # 365×144 逐日电价
assert PM.shape == (365, 144)
print(f"附件4电价: 均值 {PM.mean():.4f}, 范围 [{PM.min():.4f}, {PM.max():.4f}]")
print(f"日均价带: min {PM.mean(1).min():.4f} / max {PM.mean(1).max():.4f}; "
      f"日内价差(max-min)均值 {(PM.max(1)-PM.min(1)).mean():.4f}")
# 与附件1固定电价比较
pc = q3.price
print(f"附件1固定电价: 均值 {pc.mean():.4f}, 日内价差 {pc.max()-pc.min():.4f}")

NDAY, T, dt = 365, 144, 1.0 / 6.0
feb1 = 31
rep = slice(feb1, NDAY)
dates = q3.dates

# ================= Q4-2: 问题2策略 × 波动电价 =================
R2 = q2.run_year(prices=PM)
E_p2 = R2["plans_p"][rep].sum() * dt
C_p2 = float(R2["plan_cost"][rep].sum())
E_e2 = R2["Emerg"][rep].sum() * dt
C_e2 = float((R2["Emerg"][rep] * PM[rep] * q2.EMULT).sum() * dt)
C2 = C_p2 + C_e2
ed2 = int((R2["Emerg"][rep].sum(1) > 1e-9).sum())
assert R2["Sall"].min() >= 1200 - 1e-6 and R2["Sall"].max() <= 10800 + 1e-6
print("\n===== Q4-2 (问题2策略 × 波动电价) =====")
print(f"计划购电量 {E_p2:,.1f} kWh, 计划购电费 {C_p2:,.2f} 元")
print(f"紧急购电量 {E_e2:,.1f} kWh, 紧急购电费 {C_e2:,.2f} 元, 天数 {ed2}/334")
print(f"总费用 {C2:,.2f} 元   (恒定电价问题2: 14,546,503.97 元, 差 {C2-14546503.97:+,.2f})")

# ================= Q4-3: 问题3策略 × 波动电价 =================
R3 = q3.run_year("V4F", GAMMA=1.02, DELTA=0.98, prices=PM)
s3 = q3.settle(R3, rep, prices=PM)
ed3 = int((R3["Emerg"][rep].sum(1) > 1e-9).sum())
assert R3["Sall"].min() >= 1200 - 1e-6 and R3["Sall"].max() <= 10800 + 1e-6
viol = sum(1 for d in range(NDAY) for t in range(T)
           if R3["Emerg"][d, t] > 1e-9 and not (R3["ch_r"][d, t] < 1e-9 and
           (R3["dis_r"][d, t] > 5000 - 1e-6 or R3["Sall"][d, t + 1] < 1200 + 1e-6)))
assert viol == 0, f"紧急购电语义违规 {viol}"
print("\n===== Q4-3 (问题3策略 × 波动电价) =====")
print(f"按调整量实付 {s3['base']:,.2f} 元, 调整违约/溢价费 {s3['fee_adj']:,.2f} 元")
print(f"紧急购电费 {s3['emerg']:,.2f} 元, 紧急购电量 {s3['E_emerg']:,.1f} kWh, 天数 {ed3}/334")
print(f"调整: 上调 {s3['E_adj_up']:,.1f} kWh, 下调 {s3['E_adj_dn']:,.1f} kWh")
print(f"总费用 {s3['total']:,.2f} 元   (恒定电价问题3: 13,825,637.87 元, 差 {s3['total']-13825637.87:+,.2f})")

# ---------- 表3 指定日期 ----------
def tag(m):
    plus = "+1" if m >= 1440 else ""
    m %= 1440
    return f"{m//60}:{m%60:02d}{plus}"

blocks = [[143] + list(range(0, 23))] + [list(range(a, b))
         for a, b in [(23, 47), (47, 71), (71, 95), (95, 119), (119, 143)]]
bnames = ["0:00-4:00","4:00-8:00","8:00-12:00","12:00-16:00","16:00-20:00","20:00-24:00"]
idx_t1 = [59, 71, 83, 95, 107, 119]
names_t1 = ["10:00-10:10","12:00-12:10","14:00-14:10",
            "16:00-16:10","18:00-18:10","20:00-20:10"]

def fmt_date(d):
    x = dates[d]
    return f"{x.year}.{x.month}.{x.day}"

def emerg_windows(e):
    wins = []; t = 0
    while t < T:
        if e[t] > 1e-9:
            j = t
            while j + 1 < T and e[j + 1] > 1e-9: j += 1
            wins.append([f"{tag(10*(t+1))}-{tag(10*(j+2))}", float(e[t:j+1].sum() * dt)])
            t = j + 1
        else: t += 1
    return wins

table3 = {}
for d in range(NDAY):
    if fmt_date(d) in ["2025.3.20", "2025.6.21", "2025.9.23", "2025.12.21"]:
        entry = {"date": fmt_date(d)}
        # Q4-2
        Ep2 = R2["plans_p"][d] * dt
        entry["q42"] = {
            "plan_t1": {n: float(Ep2[i]) for n, i in zip(names_t1, idx_t1)},
            "plan_total": float(Ep2.sum()), "plan_cost": float(R2["plan_cost"][d]),
            "blocks": {n: [float((R2["ch_r"][d][b]*dt).sum()), float((R2["dis_r"][d][b]*dt).sum())]
                       for n, b in zip(bnames, blocks)},
            "s000": float(R2["Sall"][d - 1][143] if d > 0 else 6000.0),
            "s2400": float(R2["Sall"][d][143]),
            "emerg_windows": emerg_windows(R2["Emerg"][d]),
            "emerg_total": float(R2["Emerg"][d].sum() * dt)}
        # Q4-3
        Ep3, Ea3 = R3["plans_p"][d] * dt, R3["adj_p"][d] * dt
        dev = np.abs(Ea3 - Ep3)
        entry["q43"] = {
            "plan_t1": {n: float(Ep3[i]) for n, i in zip(names_t1, idx_t1)},
            "adj_t1": {n: float(Ea3[i]) for n, i in zip(names_t1, idx_t1)},
            "plan_total": float(Ep3.sum()), "adj_total": float(Ea3.sum()),
            "plan_cost": float(R3["plan_cost"][d]),
            "adj_settle_cost": float((Ea3 * PM[d]).sum() + 0.5 * (dev * PM[d]).sum()),
            "blocks": {n: [float((R3["ch_r"][d][b]*dt).sum()), float((R3["dis_r"][d][b]*dt).sum())]
                       for n, b in zip(bnames, blocks)},
            "s000": float(R3["Sall"][d - 1][143] if d > 0 else 6000.0),
            "s2400": float(R3["Sall"][d][143]),
            "emerg_windows": emerg_windows(R3["Emerg"][d]),
            "emerg_total": float(R3["Emerg"][d].sum() * dt)}
        table3[fmt_date(d)] = entry

print("\n========== 表3 指定日期(波动电价) ==========")
for k, v in table3.items():
    print(f"\n--- {k} ---")
    a, b = v["q42"], v["q43"]
    print(f"  Q4-2: 计划购电量 {a['plan_total']:.2f} kWh, 购电费 {a['plan_cost']:.2f} 元, 紧急 {a['emerg_total']:.2f} kWh")
    print(f"  Q4-3: 计划 {b['plan_total']:.2f} -> 调整 {b['adj_total']:.2f} kWh, "
          f"计划费 {b['plan_cost']:.2f} -> 结算 {b['adj_settle_cost']:.2f} 元, 紧急 {b['emerg_total']:.2f} kWh")

json.dump({"q42": {"E_plan": E_p2, "C_plan": C_p2, "E_emerg": E_e2,
                   "C_emerg": C_e2, "C_total": C2, "emerg_days": ed2},
           "q43": {**s3, "emerg_days": ed3},
           "price_stats": {"mean": float(PM.mean()), "min": float(PM.min()),
                           "max": float(PM.max()),
                           "daily_spread_mean": float((PM.max(1) - PM.min(1)).mean())},
           "table3": table3},
          open(BASE + "/Q4/summary4.json", "w"), ensure_ascii=False, indent=1)

# ---------- 写 result4-2.xlsx (同 result2 版式) ----------
import datetime as _dt

def write_result4(fname, R, has_adj):
    out = openpyxl.Workbook()
    hdr = ["日期\\时间"] + [f"{tag(10*(i+1))}-{tag(10*(i+2))}" for i in range(143)]
    hdr += ["0:00-0:10+1", "全天购电量", "全天购电费"]
    hdr[42] = "7:0-7:10"
    ws1 = out.active; ws1.title = "计划购电量"
    ws1.append(hdr)
    for d in range(feb1, NDAY):
        Ep = R["plans_p"][d] * dt
        ws1.append([dates[d]] + [round(float(x), 4) for x in Ep] +
                   [round(float(Ep.sum()), 4), round(float(R["plan_cost"][d]), 4)])
    if has_adj:
        ws1b = out.create_sheet("调整购电量")
        ws1b.append(hdr.copy())
        for d in range(feb1, NDAY):
            Ea = R["adj_p"][d] * dt
            ws1b.append([dates[d]] + [round(float(x), 4) for x in Ea] +
                        [round(float(Ea.sum()), 4), round(float((R["adj_p"][d] * PM[d]).sum() * dt), 4)])
    ws2 = out.create_sheet("充放电量")
    ws2.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for d in range(feb1, NDAY):
        Ech, Edis = R["ch_r"][d] * dt, R["dis_r"][d] * dt
        s000 = R["Sall"][d - 1][143] if d > 0 else 6000.0
        s2400 = R["Sall"][d][143]
        for k, (n, b) in enumerate(zip(bnames, blocks)):
            r = [dates[d] if k == 0 else None, n,
                 round(float(Ech[b].sum()), 4), round(float(Edis[b].sum()), 4)]
            if k == 0:   r += [_dt.time(0, 0), round(float(s000), 4)]
            elif k == 1: r += ["24:00", round(float(s2400), 4)]
            else:        r += [None, None]
            ws2.append(r)
    ws3 = out.create_sheet("紧急购电量")
    ws3.append(["日期", "购电时间段", "购电量"])
    for d in range(feb1, NDAY):
        first = True
        for w, energy in emerg_windows(R["Emerg"][d]):
            ws3.append([dates[d] if first else None, w, round(energy, 4)])
            first = False
    out.save(BASE + f"/Q4/{fname}")

write_result4("result4-2.xlsx", R2, has_adj=False)
write_result4("result4-3.xlsx", R3, has_adj=True)
print("\n已写出 result4-2.xlsx, result4-3.xlsx")

np.savez(BASE + "/Q4/q4_data.npz",
         R2_plans_p=R2["plans_p"], R2_Emerg=R2["Emerg"], R2_ch_r=R2["ch_r"],
         R2_dis_r=R2["dis_r"], R2_Sall=R2["Sall"], R2_plan_cost=R2["plan_cost"],
         R3_plans_p=R3["plans_p"], R3_adj_p=R3["adj_p"], R3_Emerg=R3["Emerg"],
         R3_ch_r=R3["ch_r"], R3_dis_r=R3["dis_r"], R3_Sall=R3["Sall"],
         R3_plan_cost=R3["plan_cost"], PM=PM, allow_pickle=True)
print("已保存 q4_data.npz")
