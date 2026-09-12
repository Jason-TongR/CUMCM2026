#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Q2 策略寻优实验: 预测方法 × 余量系数 的全年总费用对比"""
import numpy as np
from pathlib import Path
import openpyxl
from scipy.optimize import linprog

BASE = Path(__file__).resolve().parents[1]
c1 = openpyxl.load_workbook(BASE / "附件" / "附件1.xlsx", data_only=True)["Sheet1"]
rows1 = list(c1.iter_rows(min_row=2, values_only=True))
price = np.array([float(r[1]) for r in rows1])
L_typ = np.array([float(r[2]) for r in rows1])
PV_typ = np.array([float(r[3]) for r in rows1])

def load_cols(path, sheet):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return [r[0] for r in rows], np.array([[float(v) for v in r[1:145]] for r in rows])

dates, LD = load_cols(BASE / "附件" / "附件2.xlsx", "小区负载")
_, PVD = load_cols(BASE / "附件" / "附件2.xlsx", "光伏发电实际功率")
NDAY = 365
T, dt = 144, 1.0 / 6.0
ETA_C = ETA_D = 0.9
PMAX, SMIN, SMAX = 5000.0, 1200.0, 10800.0
EMULT = 5.0
ip, ich, idi, iq, is_ = 0, T, 2 * T, 3 * T, 4 * T
nx = 5 * T + 1
feb1 = 31

def plan_lp(Lhat, PVhat, S0, lscale=1.0, pscale=1.0):
    f = np.zeros(nx); f[ip:ip+T] = price * dt
    Aeq, beq = [], []
    Lh, Ph = Lhat * lscale, PVhat * pscale
    for t in range(T):
        r = np.zeros(nx)
        r[ip+t] = 1; r[idi+t] = 1; r[ich+t] = -1; r[iq+t] = -1
        Aeq.append(r); beq.append(Lh[t] - Ph[t])
    for t in range(T):
        r = np.zeros(nx)
        r[is_+t+1] = 1; r[is_+t] = -1
        r[ich+t] = -ETA_C*dt; r[idi+t] = dt/ETA_D
        Aeq.append(r); beq.append(0.0)
    for col, val in [(is_, S0), (is_+T, S0)]:
        r = np.zeros(nx); r[col] = 1
        Aeq.append(r); beq.append(val)
    bnd = ([(0, None)]*T + [(0, PMAX)]*2*T + [(0, None)]*T + [(SMIN, SMAX)]*(T+1))
    res = linprog(f, A_eq=np.array(Aeq), b_eq=np.array(beq), bounds=bnd, method="highs")
    assert res.status == 0
    x = res.x
    return x[ip:ip+T], x[ich:ich+T], x[idi:idi+T], res.fun

def simulate(ps, chs, diss, L, PV, S0):
    e = np.zeros(T)
    S = np.empty(T+1); S[0] = S0
    for t in range(T):
        dis_max = min(PMAX, max(S[t]-SMIN, 0.0)*ETA_D/dt)
        ch_max  = min(PMAX, max(SMAX-S[t], 0.0)/ETA_C/dt)
        dis = min(diss[t], dis_max); ch = min(chs[t], ch_max)
        resid = (L[t] - ps[t] - PV[t]) - (dis - ch)
        if resid > 0:
            r = min(ch, resid); ch -= r; resid -= r
            r = min(dis_max - dis, resid); dis += r; resid -= r
            e[t] = max(resid, 0.0)
        else:
            sur = -resid
            r = min(dis, sur); dis -= r; sur -= r
            r = min(ch_max - ch, sur); ch += r; sur -= r
        S[t+1] = S[t] + ETA_C*ch*dt - dis*dt/ETA_D
    return e, S

def forecast(d, method):
    if d == 0:
        return L_typ, PV_typ
    if method == "pers":
        return LD[d-1], PVD[d-1]
    if method == "w1":                       # 上周同星期
        return (LD[d-7], PVD[d-7]) if d >= 7 else (LD[d-1], PVD[d-1])
    if method.startswith("wk"):              # 最近k个同星期均值
        k = int(method[2:])
        idx = [d-7*i for i in range(1, k+1) if d-7*i >= 0]
        if not idx: return LD[d-1], PVD[d-1]
        return LD[idx].mean(0), PVD[idx].mean(0)
    if method == "ma7":
        k = max(1, d-7)
        return LD[k:d].mean(0), PVD[k:d].mean(0)
    raise ValueError(method)

def run(method, lscale=1.0, pscale=1.0, d0=0, d1=365):
    S0 = 6000.0
    Cp = np.zeros(NDAY); Ce = np.zeros(NDAY); Ee = np.zeros(NDAY)
    for d in range(NDAY):
        Lhat, PVhat = forecast(d, method)
        ps, chs, diss, pc = plan_lp(Lhat, PVhat, S0, lscale, pscale)
        e, S = simulate(ps, chs, diss, LD[d], PVD[d], S0)
        Cp[d] = pc
        Ce[d] = (e*price*EMULT).sum()*dt
        Ee[d] = e.sum()*dt
        S0 = S[T]
    sl = slice(max(d0, feb1), d1)
    return Cp[sl].sum() + Ce[sl].sum(), Cp[sl].sum(), Ce[sl].sum(), Ee[sl].sum()

import itertools, time
methods = ["pers", "w1", "wk2", "wk4", "ma7"]
scales = [(1.0, 1.0), (1.05, 1.0), (1.10, 1.0), (1.05, 0.95), (1.10, 0.90)]
print(f"{'预测':<6}{'γ负载':>6}{'δ光伏':>6}{'总费用(万元)':>14}{'计划费(万元)':>14}{'紧急费(万元)':>14}{'紧急量(MWh)':>13}")
results = {}
for m, (g, dd) in itertools.product(methods, scales):
    t0 = time.time()
    Ctot, Cp, Ce, Ee = run(m, g, dd)
    results[(m, g, dd)] = (Ctot, Cp, Ce, Ee)
    print(f"{m:<6}{g:>6.2f}{dd:>6.2f}{Ctot/1e4:>14.1f}{Cp/1e4:>14.1f}{Ce/1e4:>14.1f}{Ee/1e3:>13.1f}  ({time.time()-t0:.0f}s)", flush=True)

best = min(results.items(), key=lambda kv: kv[1][0])
print("\n最优组合:", best[0], f"总费用 {best[1][0]:,.0f} 元")

# 前后半段稳健性: 用最优参数看 Feb-Aug vs Sep-Dec
(m, g, dd), _ = best
C1 = run(m, g, dd, 31, 243)[0]   # Feb-Aug
C2 = run(m, g, dd, 243, 365)[0]  # Sep-Dec
print(f"最优参数分段: 2-8月 {C1/1e4:.1f} 万元, 9-12月 {C2/1e4:.1f} 万元")
