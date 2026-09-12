#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Q3 最终策略 V4F (γ=1.02, δ=0.98): 生成 result3.xlsx + summary3.json + q3_data.npz + 全部校验"""
import json
import numpy as np
import openpyxl
import solve_q3 as q

BASE = q.BASE
R = q.run_year("V4FR", GAMMA=1.02, DELTA=0.98, S_MIN_END=4000.0)
rep = slice(q.feb1, 365)
price, NDAY, T, dt = q.price, q.NDAY, q.T, q.dt
plans_p, adj_p = R["plans_p"], R["adj_p"]
Emerg, ch_r, dis_r, Sall = R["Emerg"], R["ch_r"], R["dis_r"], R["Sall"]

# ---------- 校验 ----------
assert Sall.min() >= 1200 - 1e-6 and Sall.max() <= 10800 + 1e-6, "储能越界"
assert (adj_p >= -1e-9).all()
up_total = float(np.maximum(adj_p - plans_p, 0).sum() * dt)
print(f"  上调电量 {up_total:,.1f} kWh (V4F 双向调整)")
viol = sum(1 for d in range(NDAY) for t in range(T)
           if Emerg[d, t] > 1e-9 and not (ch_r[d, t] < 1e-9 and
           (dis_r[d, t] > 5000 - 1e-6 or Sall[d, t + 1] < 1200 + 1e-6)))
assert viol == 0, f"紧急购电语义违规 {viol}"
s = q.settle(R, rep)
print("========== 问题3 最终策略 V4FR (标定+双向调整+S(24:00)≥4000备用, γ=1.02, δ=0.98) ==========")
print(f"校验全部通过. 年度(2.1-12.31):")
print(f"  计划购电费(0:00计划口径) {s['plan_nominal']:,.2f} 元")
print(f"  按调整量实付 {s['base']:,.2f} 元, 调整违约/溢价费 {s['fee_adj']:,.2f} 元(上下调各0.5倍)")
print(f"  紧急购电费 {s['emerg']:,.2f} 元, 紧急购电量 {s['E_emerg']:,.1f} kWh")
print(f"  调整下调电量 {s['E_adj_dn']:,.1f} kWh, 上调电量 {s['E_adj_up']:,.1f} kWh")
print(f"  总费用 {s['total']:,.2f} 元")
edays = int((Emerg[rep].sum(1) > 1e-9).sum())
print(f"  紧急购电天数 {edays}/334")

# ---------- 表3指定日期 ----------
def tag(m):
    plus = "+1" if m >= 1440 else ""
    m %= 1440
    return f"{m//60}:{m%60:02d}{plus}"

# 表2 钟点日时段块: "0:00-4:00" = 前一日末段(0:00-0:10) + 当日第1-23段
def bsum(arr, d, k):
    if k == 0:
        return float(arr[d - 1][143] + arr[d][0:23].sum())
    seg = [(23, 47), (47, 71), (71, 95), (95, 119), (119, 143)][k - 1]
    return float(arr[d][seg[0]:seg[1]].sum())
bnames = ["0:00-4:00","4:00-8:00","8:00-12:00","12:00-16:00","16:00-20:00","20:00-24:00"]
idx_t1 = [59, 71, 83, 95, 107, 119]
names_t1 = ["10:00-10:10","12:00-12:10","14:00-14:10",
            "16:00-16:10","18:00-18:10","20:00-20:10"]

def fmt_date(d):
    x = q.dates[d]
    return f"{x.year}.{x.month}.{x.day}"

def emerg_windows(e):
    wins = []; t = 0
    while t < T:
        if e[t] > 1e-9:
            j = t
            while j + 1 < T and e[j + 1] > 1e-9: j += 1
            wins.append([f"{tag(10*(t+1))}-{tag(10*(j+2))}", float(e[t:j+1].sum() * dt)])
            t = j + 1
        else: t += 1
    return wins

table3 = {}
for d in range(NDAY):
    if fmt_date(d) in ["2025.3.20", "2025.6.21", "2025.9.23", "2025.12.21"]:
        Ep, Ea = plans_p[d] * dt, adj_p[d] * dt
        Ech, Edis = ch_r[d] * dt, dis_r[d] * dt
        e = Emerg[d] * dt
        dev = np.abs(Ea - Ep)
        fee_dn = float((dev * price * 0.5).sum())
        table3[fmt_date(d)] = {
            "plan_t1": {n: float(Ep[i]) for n, i in zip(names_t1, idx_t1)},
            "adj_t1": {n: float(Ea[i]) for n, i in zip(names_t1, idx_t1)},
            "plan_total": float(Ep.sum()), "adj_total": float(Ea.sum()),
            "plan_cost": float(R["plan_cost"][d]),
            "adj_settle_cost": float((Ea * price).sum() + fee_dn),
            "blocks": {n: [bsum(ch_r * dt, d, k), bsum(dis_r * dt, d, k)]
                       for k, n in enumerate(bnames)},
            "s000": float(Sall[d - 1][143] if d > 0 else 6000.0),
            "s2400": float(Sall[d][143]),
            "emerg_windows": emerg_windows(Emerg[d]),
            "emerg_total": float(e.sum()),
        }

print("\n========== 表3 指定日期 ==========")
for k, v in table3.items():
    print(f"\n--- {k} ---")
    print(f"  计划购电量 {v['plan_total']:.2f} kWh -> 调整购电量 {v['adj_total']:.2f} kWh")
    print(f"  计划购电费 {v['plan_cost']:.2f} 元, 调整后结算(含违约费) {v['adj_settle_cost']:.2f} 元")
    print("  表1(调整后):", {n: round(x, 2) for n, x in v["adj_t1"].items()})
    print("  表2:", {n: [round(a, 2), round(b, 2)] for n, (a, b) in v["blocks"].items()})
    print(f"  s(0:00)={v['s000']:.2f}, s(24:00)={v['s2400']:.2f}")
    print(f"  紧急购电: {[(w, round(x,2)) for w, x in v['emerg_windows']] if v['emerg_windows'] else '无'} (合计 {v['emerg_total']:.2f} kWh)")

json.dump({"annual": s, "emerg_days": edays, "table3": table3},
          open(BASE / "Q3" / "summary3.json", "w"), ensure_ascii=False, indent=1)

# ---------- 写 result3.xlsx ----------
out = openpyxl.Workbook()
hdr = ["日期\\时间"] + [f"{tag(10*(i+1))}-{tag(10*(i+2))}" for i in range(143)]
hdr += ["0:00-0:10+1", "全天购电量", "全天购电费"]
hdr[42] = "7:0-7:10"          # 保留模板 typo

ws1 = out.active; ws1.title = "计划购电量"
ws1.append(hdr)
for d in range(q.feb1, NDAY):
    Ep = plans_p[d] * dt
    ws1.append([q.dates[d]] + [round(float(x), 4) for x in Ep] +
               [round(float(Ep.sum()), 4), round(float(R["plan_cost"][d]), 4)])

ws1b = out.create_sheet("调整购电量")
ws1b.append(hdr.copy())
for d in range(q.feb1, NDAY):
    Ea = adj_p[d] * dt
    ws1b.append([q.dates[d]] + [round(float(x), 4) for x in Ea] +
                [round(float(Ea.sum()), 4), round(float((adj_p[d] * price).sum() * dt), 4)])

import datetime as _dt
ws2 = out.create_sheet("充放电量")
ws2.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
for d in range(q.feb1, NDAY):
    Ech, Edis = ch_r[d] * dt, dis_r[d] * dt
    s000 = Sall[d - 1][143] if d > 0 else 6000.0
    s2400 = Sall[d][143]
    for k, n in enumerate(bnames):
        r = [q.dates[d] if k == 0 else None, n,
             round(bsum(ch_r * dt, d, k), 4), round(bsum(dis_r * dt, d, k), 4)]
        if k == 0:   r += [_dt.time(0, 0), round(float(s000), 4)]
        elif k == 1: r += ["24:00", round(float(s2400), 4)]
        else:        r += [None, None]
        ws2.append(r)

ws3 = out.create_sheet("紧急购电量")
ws3.append(["日期", "购电时间段", "购电量"])
for d in range(q.feb1, NDAY):
    first = True
    for w, energy in emerg_windows(Emerg[d]):
        ws3.append([q.dates[d] if first else None, w, round(energy, 4)])
        first = False
out.save(BASE / "Q3" / "result3.xlsx")
print("\n已写出 result3.xlsx")

np.savez(BASE / "Q3" / "q3_data.npz", plans_p=plans_p, adj_p=adj_p, Emerg=Emerg,
         ch_r=ch_r, dis_r=dis_r, Sall=Sall, plan_cost=R["plan_cost"], allow_pickle=True)
print("已保存 q3_data.npz")
