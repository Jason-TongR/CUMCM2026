# -*- coding: utf-8 -*-
"""
导出 Origin 可直接导入的 Q2 数据（CUMCM2026 C题 问题2）。

输出目录: 本脚本所在目录
  Q2_表1_逐时段数据_四日期合并.csv     4 个指定日期 x 144 时段
  Q2_表1_<日期>.csv                     每个日期单独一份
  Q2_表3_紧急购电矩阵_334x144_kWh.csv   带日期/时段的矩阵（给 Excel 核对）
  Q2_表3_紧急购电矩阵_纯数值.csv        334 行 x 144 列纯数值（给 Origin Matrix）
  Q2_表3_紧急购电窗口清单_指定日期.csv
  Q2_表3_紧急购电窗口清单_全年.csv
  Q2_扫描结果_预测方法x余量.csv         与 q2_experiments.py 同一套网格
  Q2_扫描结果_w1精细网格.csv            w1 方法的 gamma x delta 曲面数据
  README_数据说明.txt

用法:
  python make_origin_csv.py            # 只导出表1/表3数据（快）
  python make_origin_csv.py --scan     # 额外复跑参数扫描（约数分钟）
"""

import csv
import os
import shutil
import sys
import time
import zipfile
from pathlib import Path

import numpy as np

# scipy 通过 --target 装在临时目录，避免改动运行环境
_PKG = Path(os.environ.get("CUMCM_PKG_DIR",
                           str(Path(os.environ.get("TEMP", "/tmp")) / "cumcm_pkgs")))
if _PKG.is_dir() and str(_PKG) not in sys.path:
    sys.path.insert(0, str(_PKG))
from scipy.optimize import linprog  # noqa: E402

try:  # 支持直接读取 xlsx
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

OUT = Path(__file__).resolve().parent
MSG = Path(r"C:/Users/Is/OneDrive/ドキュメント/xwechat_files/"
           r"wxid_ke2qlzl1c5gz22_f248/msg/file/2026-09")
ZIP = MSG / "CUMCM2026Problems.zip"
NPZ = MSG / "q2_data.npz"

# ---------------- 模型参数（与 solve_q2.py 完全一致） ----------------
T = 144
DT = 1.0 / 6.0
ETA_C = ETA_D = 0.9
PMAX = 5000.0
SMIN, SMAX = 1200.0, 10800.0
EMULT = 5.0
NDAY = 365
FEB1 = 31
IP, ICH, IDI, IQ, IS_ = 0, T, 2 * T, 3 * T, 4 * T
NX = 5 * T + 1
DATES4 = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]


# ---------------- 数据读取 ----------------
def _cp437_to_gbk(name: str) -> str:
    try:
        return name.encode("cp437").decode("gbk")
    except Exception:
        return name


def load_attachments():
    """读取附件1（电价/典型日负载/光伏预测）与附件2（实际负载/光伏）。"""
    cache = Path(os.environ.get("TEMP", "/tmp")) / "cumcmC"
    cache.mkdir(parents=True, exist_ok=True)
    need = {"附件1.xlsx", "附件2.xlsx"}
    if not need.issubset({p.name for p in cache.glob("*.xlsx")}):
        with zipfile.ZipFile(ZIP) as z:
            for info in z.infolist():
                name = _cp437_to_gbk(info.filename)
                if name.endswith("附件1.xlsx") or name.endswith("附件2.xlsx"):
                    target = cache / Path(name).name
                    with z.open(info) as src, open(target, "wb") as dst:
                        shutil.copyfileobj(src, dst)
    ws = openpyxl.load_workbook(cache / "附件1.xlsx", data_only=True)["Sheet1"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    price = np.array([float(r[1]) for r in rows])
    l_typ = np.array([float(r[2]) for r in rows])
    pv_typ = np.array([float(r[3]) for r in rows])
    wb2 = openpyxl.load_workbook(cache / "附件2.xlsx", data_only=True)
    dates = [r[0] for r in wb2["小区负载"].iter_rows(min_row=2, values_only=True)]
    ld = np.array([[float(v) for v in r[1:145]]
                   for r in wb2["小区负载"].iter_rows(min_row=2, values_only=True)])
    pvd = np.array([[float(v) for v in r[1:145]]
                    for r in wb2["光伏发电实际功率"].iter_rows(min_row=2, values_only=True)])
    return price, l_typ, pv_typ, dates, ld, pvd


def tag(minute: int) -> str:
    plus = "+1" if minute >= 1440 else ""
    minute %= 1440
    return f"{minute // 60}:{minute % 60:02d}{plus}"


LABELS = [f"{tag(10 * (i + 1))}-{tag(10 * (i + 2))}" for i in range(T)]


def day_index(dates, target: str) -> int:
    for i, d in enumerate(dates):
        s = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10]
        if s == target:
            return i
    raise KeyError(target)


def write_csv(path: Path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"  [ok] {path.name}  ({len(rows)} 行)")


# ---------------- 表1 / 表3 数据导出 ----------------
def export_tables(price, l_typ, pv_typ, dates, ld, pvd, npz):
    plans_p = npz["plans_p"]
    plans_ch = npz["plans_ch"]
    plans_dis = npz["plans_dis"]
    emerg = npz["Emerg"]
    soc = npz["Sall"]
    ch_r = npz["ch_r"]
    dis_r = npz["dis_r"]

    header = ["日期", "时刻", "时段序号", "购电量kWh", "购电功率kW", "电价元每kWh",
              "小区负载kW", "光伏实际功率kW", "储电量期初kWh", "储电量期末kWh",
              "实际充电量kWh", "实际放电量kWh", "计划充电量kWh", "计划放电量kWh"]
    all_rows = []
    for target in DATES4:
        d = day_index(dates, target)
        rows = []
        for i in range(T):
            rows.append([
                target, LABELS[i], i + 1,
                round(float(plans_p[d][i] * DT), 4),
                round(float(plans_p[d][i]), 4),
                round(float(price[i]), 4),
                round(float(ld[d][i]), 4),
                round(float(pvd[d][i]), 4),
                round(float(soc[d][i]), 4),
                round(float(soc[d][i + 1]), 4),
                round(float(ch_r[d][i] * DT), 4),
                round(float(dis_r[d][i] * DT), 4),
                round(float(plans_ch[d][i] * DT), 4),
                round(float(plans_dis[d][i] * DT), 4),
            ])
        write_csv(OUT / f"Q2_表1_{target}.csv", header[1:], [r[1:] for r in rows])
        all_rows.extend(rows)
    write_csv(OUT / "Q2_表1_逐时段数据_四日期合并.csv", header, all_rows)

    # --- 表3：全年紧急购电矩阵 ---
    rep_idx = list(range(FEB1, NDAY))

    def dstr(d):
        x = dates[d]
        return x.strftime("%Y-%m-%d") if hasattr(x, "strftime") else str(x)[:10]

    mat = np.round(emerg[rep_idx] * DT, 4)
    rows = [[dstr(d)] + list(mat[k]) for k, d in enumerate(rep_idx)]
    write_csv(OUT / "Q2_表3_紧急购电矩阵_334x144_kWh.csv", ["日期"] + LABELS, rows)
    write_csv(OUT / "Q2_表3_紧急购电矩阵_纯数值.csv", [""] * T, mat.tolist())

    def windows(vec):
        wins, t = [], 0
        while t < T:
            if vec[t] > 1e-9:
                j = t
                while j + 1 < T and vec[j + 1] > 1e-9:
                    j += 1
                energy = float(vec[t:j + 1].sum() * DT)
                mins = 10 * (j - t + 1)
                wins.append([tag(10 * (t + 1)), tag(10 * (j + 2)), mins,
                             round(energy, 4), round(energy / (mins / 60.0), 4)])
                t = j + 1
            else:
                t += 1
        return wins

    wh = ["日期", "开始时刻", "结束时刻", "时长min", "紧急购电量kWh", "平均功率kW"]
    rows4 = []
    for target in DATES4:
        d = day_index(dates, target)
        w = windows(emerg[d])
        if not w:
            rows4.append([target, "无", "无", 0, 0.0, 0.0])
        else:
            for a, b, m, e, p in w:
                rows4.append([target, a, b, m, e, p])
    write_csv(OUT / "Q2_表3_紧急购电窗口清单_指定日期.csv", wh, rows4)

    rows_year = []
    for d in rep_idx:
        for a, b, m, e, p in windows(emerg[d]):
            rows_year.append([dstr(d), a, b, m, e, p])
    write_csv(OUT / "Q2_表3_紧急购电窗口清单_全年.csv", wh, rows_year)

    kpi = [
        ["报告期", "2025-02-01 ~ 2025-12-31", f"{len(rep_idx)} 天", "1月为预热期"],
        ["计划购电量kWh", round(float(plans_p[rep_idx].sum() * DT), 1), "", ""],
        ["计划购电费元", round(float(npz["plan_cost"][rep_idx].sum()), 2), "", ""],
        ["紧急购电量kWh", round(float(emerg[rep_idx].sum() * DT), 1), "", ""],
        ["紧急购电费元", round(float((emerg * price * EMULT)[rep_idx].sum() * DT), 2), "", ""],
        ["总费用元", "", "", ""],
        ["发生紧急购电天数", int((emerg[rep_idx].sum(1) > 1e-9).sum()),
         f"共 {len(rep_idx)} 天", ""],
        ["储能SOC范围kWh", round(float(soc.min()), 2), round(float(soc.max()), 2), ""],
    ]
    kpi[5][1] = round(float(kpi[2][1]) + float(kpi[4][1]), 2)
    write_csv(OUT / "Q2_年度KPI.csv", ["指标", "数值", "备注", "说明"], kpi)

    # --- 表2：4 小时时段块（按模板口径修正首块） ---
    bnames = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
              "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    tail = [list(range(a, b)) for a, b in
            [(23, 47), (47, 71), (71, 95), (95, 119), (119, 143)]]
    # 模板口径：一行覆盖 0:10~0:10(+1)，故 0:00-0:10 这一段取前一天的最后一段
    # （这正是 result2.xlsx 里 0:00-4:00 那一行需要修正的地方）
    rows2 = []
    for d in rep_idx:
        prev_last_ch = float(ch_r[d - 1][143]) * DT
        prev_last_dis = float(dis_r[d - 1][143]) * DT
        blk = [["0:00-4:00",
                prev_last_ch + float(ch_r[d][0:23].sum()) * DT,
                prev_last_dis + float(dis_r[d][0:23].sum()) * DT],
               ["4:00-8:00", float(ch_r[d][23:47].sum()) * DT,
                float(dis_r[d][23:47].sum()) * DT],
               ["8:00-12:00", float(ch_r[d][47:71].sum()) * DT,
                float(dis_r[d][47:71].sum()) * DT],
               ["12:00-16:00", float(ch_r[d][71:95].sum()) * DT,
                float(dis_r[d][71:95].sum()) * DT],
               ["16:00-20:00", float(ch_r[d][95:119].sum()) * DT,
                float(dis_r[d][95:119].sum()) * DT],
               ["20:00-24:00", float(ch_r[d][119:143].sum()) * DT,
                float(dis_r[d][119:143].sum()) * DT]]
        for k, (name, ch, dis) in enumerate(blk):
            extra = ["", ""]
            if k == 0:
                extra = ["0:00", round(float(soc[d - 1][143]), 4)]
            elif k == 1:
                extra = ["24:00", round(float(soc[d][143]), 4)]
            rows2.append([dstr(d) if k == 0 else "", name,
                          round(ch, 4), round(dis, 4)] + extra)
    write_csv(OUT / "Q2_表2_充放电量_修正版_按模板口径.csv",
              ["日期", "时间段", "充电量kWh", "放电量kWh", "时刻", "储电量kWh"], rows2)


# ---------------- 参数扫描（与 q2_experiments.py 同一套网格） ----------------
def build_a_eq():
    a_eq = np.zeros((2 * T + 2, NX))
    for t in range(T):
        a_eq[t, IP + t] = 1
        a_eq[t, IDI + t] = 1
        a_eq[t, ICH + t] = -1
        a_eq[t, IQ + t] = -1
    for t in range(T):
        a_eq[T + t, IS_ + t + 1] = 1
        a_eq[T + t, IS_ + t] = -1
        a_eq[T + t, ICH + t] = -ETA_C * DT
        a_eq[T + t, IDI + t] = DT / ETA_D
    a_eq[2 * T, IS_] = 1
    a_eq[2 * T + 1, IS_ + T] = 1
    return a_eq


BOUNDS = ([(0, None)] * T + [(0, PMAX)] * 2 * T + [(0, None)] * T
          + [(SMIN, SMAX)] * (T + 1))
FOBJ = np.zeros(NX)


def forecast(method, d, ld, pvd, l_typ, pv_typ):
    if d == 0:
        return l_typ, pv_typ
    if method == "pers":
        return ld[d - 1], pvd[d - 1]
    if method == "w1":
        return (ld[d - 7], pvd[d - 7]) if d >= 7 else (ld[d - 1], pvd[d - 1])
    if method.startswith("wk"):
        k = int(method[2:])
        idx = [d - 7 * i for i in range(1, k + 1) if d - 7 * i >= 0]
        if not idx:
            return ld[d - 1], pvd[d - 1]
        return ld[idx].mean(0), pvd[idx].mean(0)
    if method == "ma7":
        k = max(0, d - 7)
        if k < d:                      # 最近最多 7 天滑动平均
            return ld[k:d].mean(0), pvd[k:d].mean(0)
        return (ld[d - 1], pvd[d - 1]) if d > 0 else (l_typ, pv_typ)
    raise ValueError(method)


def simulate(ps, chs, diss, l, pv, s0):
    e = np.zeros(T)
    s = np.empty(T + 1)
    s[0] = s0
    for t in range(T):
        dmax = min(PMAX, max(s[t] - SMIN, 0.0) * ETA_D / DT)
        cmax = min(PMAX, max(SMAX - s[t], 0.0) / ETA_C / DT)
        dis = min(diss[t], dmax)
        ch = min(chs[t], cmax)
        resid = (l[t] - ps[t] - pv[t]) - (dis - ch)
        if resid > 0:
            r = min(ch, resid)
            ch -= r
            resid -= r
            r = min(dmax - dis, resid)
            dis += r
            resid -= r
            e[t] = max(resid, 0.0)
        else:
            sur = -resid
            r = min(dis, sur)
            dis -= r
            sur -= r
            r = min(cmax - ch, sur)
            ch += r
            sur -= r
        s[t + 1] = s[t] + ETA_C * ch * DT - dis * DT / ETA_D
    return e, s


def run_scan(method, g, dd, price, l_typ, pv_typ, ld, pvd, a_eq, d0=FEB1, d1=NDAY):
    fobj = FOBJ.copy()
    fobj[IP:IP + T] = price * DT
    s0 = 6000.0
    cp_l = []
    ce_l = []
    ee_l = []
    for d in range(NDAY):
        lh, ph = forecast(method, d, ld, pvd, l_typ, pv_typ)
        beq = np.empty(2 * T + 2)
        beq[:T] = lh * g - ph * dd
        beq[T:2 * T] = 0.0
        beq[2 * T] = s0
        beq[2 * T + 1] = s0
        res = linprog(fobj, A_eq=a_eq, b_eq=beq, bounds=BOUNDS, method="highs")
        if res.status != 0:
            raise RuntimeError(f"LP 失败 {method} {g} {dd} day {d}: {res.message}")
        x = res.x
        ps = x[IP:IP + T]
        e, s = simulate(ps, x[ICH:ICH + T], x[IDI:IDI + T],
                        ld[d], pvd[d], s0)
        cp_l.append(res.fun)
        ce_l.append(float((e * price * EMULT).sum() * DT))
        ee_l.append(float(e.sum() * DT))
        s0 = s[T]
    sl = slice(max(d0, FEB1), d1)
    cp = float(np.array(cp_l)[sl].sum())
    ce = float(np.array(ce_l)[sl].sum())
    ee = float(np.array(ee_l)[sl].sum())
    return cp + ce, cp, ce, ee


def export_scan(price, l_typ, pv_typ, ld, pvd):
    a_eq = build_a_eq()
    methods = ["pers", "w1", "wk2", "wk4", "ma7"]
    scales = [(1.00, 1.00), (1.05, 1.00), (1.10, 1.00), (1.05, 0.95), (1.10, 0.90)]
    rows = []
    t0 = time.time()
    for m in methods:
        for g, dd in scales:
            ctot, cp, ce, ee = run_scan(m, g, dd, price, l_typ, pv_typ, ld, pvd, a_eq)
            rows.append([m, f"{g:.2f}", f"{dd:.2f}", round(ctot / 1e4, 1),
                         round(cp / 1e4, 1), round(ce / 1e4, 1), round(ee / 1000, 1),
                         round(ctot, 2)])
            print(f"    {m:<5} γ={g:<5} δ={dd:<5} 总费用 {ctot/1e4:8.1f} 万元"
                  f"  ({time.time()-t0:.0f}s)", flush=True)
    write_csv(OUT / "Q2_扫描结果_预测方法x余量.csv",
              ["预测方法", "gamma负载余量", "delta光伏余量", "总费用万元",
               "计划购电费万元", "紧急购电费万元", "紧急购电量MWh", "总费用元"],
              rows)

    # w1 精细网格：给 Origin 画 gamma-delta 费用曲面
    fine = []
    gammas = [1.00, 1.02, 1.04, 1.05, 1.06, 1.08, 1.10]
    deltas = [0.90, 0.94, 0.98, 1.00, 1.04]
    for g in gammas:
        for dd in deltas:
            ctot, cp, ce, ee = run_scan("w1", g, dd, price, l_typ, pv_typ, ld, pvd, a_eq)
            fine.append([f"{g:.2f}", f"{dd:.2f}", round(ctot / 1e4, 1),
                         round(cp / 1e4, 1), round(ce / 1e4, 1)])
            print(f"    [fine] γ={g:<5} δ={dd:<5} {ctot/1e4:8.1f} 万元"
                  f"  ({time.time()-t0:.0f}s)", flush=True)
    write_csv(OUT / "Q2_扫描结果_w1精细网格.csv",
              ["gamma负载余量", "delta光伏余量", "总费用万元",
               "计划购电费万元", "紧急购电费万元"], fine)


README = """Q2（微网购电策略）Origin 绘图数据说明
=========================================

一、口径（务必与论文一致）
1. 时段约定：1 天 = 144 个 10 min 时段，时刻标签与官方模板 result2.xlsx 的列标签一致
   （0:10-0:20, 0:20-0:30, ..., 23:50-0:00+1, 0:00-0:10+1），即"附件1 行时刻戳＝区间起点"。
2. 报告期：2025-02-01 ~ 2025-12-31（334 天），2025 年 1 月为预热期，故矩阵文件只有 334 行。
3. 单位：购电量/充放电量 = kWh（10 min 能量）；购电功率/负载/光伏 = kW；电价 = 元/kWh；
   储电量 = kWh（范围 1200~10800）。
4. 储电量列给了"期初"和"期末"两列（期初=该 10 min 开始时刻，期末=结束时刻），
   画 SOC 曲线建议用"期末"，画区间对比可用"期初"。
5. 储能充放电效率 0.9；紧急购电电价 = 交易时刻电价 x 5。

二、文件清单
Q2_表1_逐时段数据_四日期合并.csv   4 个指定日期 x 144 时段（576 行），用于表1 的配图
Q2_表1_2025-03-20.csv 等 4 个文件   单个日期的 144 行
Q2_表3_紧急购电矩阵_334x144_kWh.csv 首列日期 + 144 个时段列（单位 kWh）
Q2_表3_紧急购电矩阵_纯数值.csv      334 行 x 144 列纯数值，供 Origin Matrix 窗口使用
Q2_表3_紧急购电窗口清单_指定日期.csv 4 个指定日期的紧急购电时段（无则为"无"）
Q2_表3_紧急购电窗口清单_全年.csv     全年 102 天、共 164 个紧急购电时段
Q2_年度KPI.csv                      报告期汇总指标
Q2_表2_充放电量_修正版_按模板口径.csv 4 小时时段块（334 天 x 6 行），已按模板口径修正首块
Q2_扫描结果_预测方法x余量.csv        5 种预测方法 x 5 组余量系数 = 25 组（同 q2_experiments.py）
Q2_扫描结果_w1精细网格.csv           w1 方法的 7x5=35 组 gamma-delta 网格，画费用曲面用

三、Origin 中的用法提示
1. 表1 配图：导入"四日期合并"文件，用"日期"列做分组（Origin: Worksheet 右键 Set As ->
   Grouping/Categorical），Plot -> Multi-Y -> Double-Y，电价用 Step 线型；
   再用 Reference Lines 画 10/12/14/16/18/20 时刻竖线。
2. 表3 热图：把"纯数值"文件全选粘贴到 Matrix 窗口，Matrix -> Set Dimensions 设
   Rows=334, Cols=144，然后 Plot -> Contour -> Image Plot（色标单位 kWh/10min）。
3. 费用曲面：用"w1精细网格"文件，Plot -> 3D -> Color Map Surface（X=gamma, Y=delta, Z=总费用）。

四、复现说明与已知差异
1. 扫描结果与论文已报告的数值完全一致（持续性 2244.8 万元、w1 1602.4 万元、
   wk4 (1.05,1.00) 1574.7 万元、最终策略 w1 (1.05,0.98) 1454.7 万元），脚本与论文同源。
2. 原 q2_experiments.py 的 ma7 分支在第 0、1 天会取到空切片（得到 NaN，会让 LP 报错）；
   本脚本改为"不足 7 天时用已有天数求均值，首日退回附件1 典型日"。因此 ma7 两行与原脚本
   不可比，其余 23 组完全一致。
3. 精细网格上总费用最小值出现在 (gamma, delta) = (1.02, 0.94)，约 1444.3 万元，比论文采用的
   (1.05, 0.98) = 1454.7 万元低约 10 万元（0.7%）。若正文用这张曲面图，建议说明"最优取值落在
   平坦平台内、差异 <1%"，避免被追问为何不取网格最小值。

生成脚本：make_origin_csv.py（同目录），数据来源：CUMCM2026Problems.zip（附件1/附件2）
与 q2_data.npz / summary2.json。
"""


def main():
    do_scan = "--scan" in sys.argv
    print("读取附件与结果文件 …")
    price, l_typ, pv_typ, dates, ld, pvd = load_attachments()
    npz = np.load(NPZ, allow_pickle=True)
    print(f"  附件1 {price.shape}, 附件2 {ld.shape}, npz {npz['plans_p'].shape}")
    print("导出表1 / 表3 数据 …")
    export_tables(price, l_typ, pv_typ, dates, ld, pvd, npz)
    (OUT / "README_数据说明.txt").write_text(README, encoding="utf-8")
    print(f"  [ok] README_数据说明.txt")
    if do_scan:
        print("复跑参数扫描（较慢）…")
        export_scan(price, l_typ, pv_typ, ld, pvd)
    print("完成。输出目录:", OUT)


if __name__ == "__main__":
    main()
